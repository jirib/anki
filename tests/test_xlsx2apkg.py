import json
import sqlite3
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

SCRIPT = Path(__file__).resolve().parents[1] / "xlsx2apkg.py"


def run_script(args):
    return subprocess.run(
        [sys.executable, str(SCRIPT), *args],
        capture_output=True,
        text=True,
    )


def read_apkg(apkg_path, tmp_path):
    with zipfile.ZipFile(apkg_path) as z:
        assert "collection.anki2" in z.namelist()
        db_bytes = z.read("collection.anki2")

    db_path = tmp_path / "collection.anki2"
    db_path.write_bytes(db_bytes)

    conn = sqlite3.connect(db_path)
    models = json.loads(conn.execute("select models from col").fetchone()[0])
    notes = [row[0] for row in conn.execute("select flds from notes").fetchall()]
    conn.close()
    return models, notes


def write_xlsx(path, rows):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(path)


def test_build_apkg_default(tmp_path):
    inp = tmp_path / "in.xlsx"
    write_xlsx(inp, [["Front", "Back"], ["hello", "ahoj"], ["world", "svet"]])
    out = tmp_path / "out.apkg"

    result = run_script([str(inp), "--output", str(out)])
    assert result.returncode == 0, result.stderr

    models, notes = read_apkg(out, tmp_path)
    assert len(models) == 1
    model = next(iter(models.values()))
    assert [f["name"] for f in model["flds"]] == ["Front", "Back"]
    assert [(t["name"], t["qfmt"], t["afmt"]) for t in model["tmpls"]] == [
        ("Card 1", "{{Front}}", '{{Front}}<hr id="answer">{{Back}}')
    ]
    assert notes == ["hello\x1fahoj", "world\x1fsvet"]


def test_build_apkg_custom_template(tmp_path):
    inp = tmp_path / "in.xlsx"
    write_xlsx(inp, [["Q", "A"], ["one", "two"]])
    template = tmp_path / "template.json"
    template.write_text(
        json.dumps(
            {"name": "Reverse", "qfmt": "{{A}}", "afmt": '{{A}}<hr id="answer">{{Q}}'}
        ),
        encoding="utf-8",
    )
    out = tmp_path / "out.apkg"

    result = run_script([str(inp), "--output", str(out), "--template", str(template)])
    assert result.returncode == 0, result.stderr

    models, notes = read_apkg(out, tmp_path)
    model = next(iter(models.values()))
    assert [(t["name"], t["qfmt"], t["afmt"]) for t in model["tmpls"]] == [
        ("Reverse", "{{A}}", '{{A}}<hr id="answer">{{Q}}')
    ]
    assert notes == ["one\x1ftwo"]


def test_build_apkg_sheet_selection(tmp_path):
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Ignore"
    ws1.append(["Front", "Back"])
    ws1.append(["nope", "nope"])
    ws2 = wb.create_sheet("Real")
    ws2.append(["Front", "Back"])
    ws2.append(["foo", "bar"])
    inp = tmp_path / "multi.xlsx"
    wb.save(inp)

    out = tmp_path / "out.apkg"
    result = run_script([str(inp), "--sheet", "Real", "--output", str(out)])
    assert result.returncode == 0, result.stderr

    models, notes = read_apkg(out, tmp_path)
    assert notes == ["foo\x1fbar"]


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__]))

"""Build an Anki .apkg deck from an Excel (.xlsx) workbook."""

import argparse

from openpyxl import load_workbook

from ankiscripts.core import build_apkg, list_templates, list_templates_with_paths


def parse_args(argv=None):
    parser = argparse.ArgumentParser(
        description="Build an Anki .apkg from an Excel (.xlsx) workbook"
    )
    parser.add_argument("input", nargs="?", help="Input .xlsx file")
    parser.add_argument("--deckname", default="My Deck", help="Deck name")
    parser.add_argument("--output", default="output.apkg", help="Output .apkg file")
    parser.add_argument("--sheet", help="Worksheet name (default: first worksheet)")
    parser.add_argument(
        "--template",
        help=(
            "Card template: a path to a JSON file or a bundled template name "
            f"({', '.join(list_templates())}). Defaults to a Front/Back card."
        ),
    )
    parser.add_argument(
        "--list-templates",
        action="store_true",
        help="List bundled template names and their paths, then exit.",
    )
    args = parser.parse_args(argv)
    if not args.list_templates and not args.input:
        parser.error("the following arguments are required: input")
    return args


def read_xlsx(path, sheet=None):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        raise ValueError(f"No rows found in {path}")
    field_names = ["" if h is None else str(h).strip() for h in header]
    data_rows = [
        ["" if c is None else str(c) for c in row]
        for row in rows
        if any(c is not None for c in row)
    ]
    wb.close()
    return field_names, data_rows


def main(argv=None):
    args = parse_args(argv)

    if args.list_templates:
        for name, path in list_templates_with_paths():
            print(f"{name}\t{path}")
        return 0

    field_names, rows = read_xlsx(args.input, args.sheet)

    build_apkg(
        field_names,
        rows,
        deckname=args.deckname,
        output=args.output,
        model_name="xlsx2anki_py",
        template=args.template,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
